import os
from pptx import Presentation

pptx_path = r"docs\RailSahayak_Stakeholder_Presentation.pptx"
pdf_path = r"docs\RailSahayak_Stakeholder_Presentation.pdf"

print("="*60)
print("PPTX Verification:")
print("File:", pptx_path)
print("File Size:", os.path.getsize(pptx_path), "bytes")
prs = Presentation(pptx_path)
print("Slide Count:", len(prs.slides))
for idx, slide in enumerate(prs.slides):
    shapes_count = len(slide.shapes)
    texts = [shape.text_frame.text[:40].replace('\n', ' ') for shape in slide.shapes if shape.has_text_frame]
    first_text = texts[0] if texts else "No text"
    print(f"  Slide {idx+1:2d}: {shapes_count:2d} shapes | Title: {first_text}")

print("\n" + "="*60)
print("PDF Verification:")
print("File:", pdf_path)
print("File Size:", os.path.getsize(pdf_path), "bytes")
with open(pdf_path, 'rb') as f:
    header = f.read(10)
    print("PDF Header:", header.decode('latin-1', errors='ignore').strip())
    f.seek(-1024, os.SEEK_END)
    tail = f.read().decode('latin-1', errors='ignore')
    print("PDF Valid EOF Marker:", '%%EOF' in tail)

with open(pdf_path, 'rb') as f:
    content = f.read().decode('latin-1', errors='ignore')
    # Count page objects in PDF
    page_count = content.count('/Type /Page\n') + content.count('/Type /Page ') + content.count('/Type/Page')
    print("PDF Page Count:", page_count)
print("="*60)
